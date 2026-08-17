"""KARNEX GST Tax Invoice — totals, amount-in-words, numbering, Excel, PDF.

Backend is the source of truth for numbers on the Tax Invoice PDF.
Seller PAN/GSTIN/bank are hardcoded constants (never env-overridden here).
"""
from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, Field

# ---------------------------------------------------------------------------
# Fixed seller constants (exact)
# ---------------------------------------------------------------------------

SELLER_NAME = "Karnex Software Solutions Private Limited"
SELLER_ADDRESS = (
    "103, Pride Purple Accord, Opp. RMZ Icon\n"
    "Near Mahalaxwiwar Hotel, Baner\n"
    "Pune, MH, India - 411045"
)
SELLER_STATE_NAME = "Maharashtra"
SELLER_STATE_CODE = "27"
SELLER_EMAIL = "karnex.singh@karnex.in"
SELLER_CIN = "U72900RJ2019OPC63826"
SELLER_PAN = "AAHCK4749A"
SELLER_GSTIN = "27AAHCK4749A1ZL"
BANK_NAME = "HDFC Bank, Baner"
BANK_NAME_SHORT = "HDFC Bank"
BANK_ACC = "50200073368143"
BANK_IFSC = "HDFC0001794"
BANK_BRANCH = "Baner, Pune"
CONTACT_WEB = "www.karnex.in"
CONTACT_EMAIL = "info@karnex.in"
CONTACT_PHONE = "+91 20 1234 5678"
CONTACT_LOC = "Pune, Maharashtra, India"

DEFAULT_INVOICE_NO = "KRSW26-27-65-VS"
DEFAULT_SAC = "998314"
DEFAULT_FOOTER = (
    "Certified that the particulars above are true and correct. "
    "The amount charged is the actual price with no additional consideration "
    "from the Service Recipient."
)

STATIC_DIR = Path(__file__).resolve().parent.parent / "static" / "tax_invoice"
LOGO_PATH = STATIC_DIR / "karnex-logo-invoice.png"
SEAL_PATH = STATIC_DIR / "karnex-seal-sign.png"

# Codes 25 and 28 intentionally absent (GST).
INDIA_STATES: list[dict[str, str]] = [
    {"code": "01", "name": "Jammu & Kashmir"},
    {"code": "02", "name": "Himachal Pradesh"},
    {"code": "03", "name": "Punjab"},
    {"code": "04", "name": "Chandigarh"},
    {"code": "05", "name": "Uttarakhand"},
    {"code": "06", "name": "Haryana"},
    {"code": "07", "name": "Delhi"},
    {"code": "08", "name": "Rajasthan"},
    {"code": "09", "name": "Uttar Pradesh"},
    {"code": "10", "name": "Bihar"},
    {"code": "11", "name": "Sikkim"},
    {"code": "12", "name": "Arunachal Pradesh"},
    {"code": "13", "name": "Nagaland"},
    {"code": "14", "name": "Manipur"},
    {"code": "15", "name": "Mizoram"},
    {"code": "16", "name": "Tripura"},
    {"code": "17", "name": "Meghalaya"},
    {"code": "18", "name": "Assam"},
    {"code": "19", "name": "West Bengal"},
    {"code": "20", "name": "Jharkhand"},
    {"code": "21", "name": "Odisha"},
    {"code": "22", "name": "Chhattisgarh"},
    {"code": "23", "name": "Madhya Pradesh"},
    {"code": "24", "name": "Gujarat"},
    {"code": "26", "name": "Dadra and Nagar Haveli and Daman and Diu"},
    {"code": "27", "name": "Maharashtra"},
    {"code": "29", "name": "Karnataka"},
    {"code": "30", "name": "Goa"},
    {"code": "31", "name": "Lakshadweep"},
    {"code": "32", "name": "Kerala"},
    {"code": "33", "name": "Tamil Nadu"},
    {"code": "34", "name": "Puducherry"},
    {"code": "35", "name": "Andaman and Nicobar Islands"},
    {"code": "36", "name": "Telangana"},
    {"code": "37", "name": "Andhra Pradesh"},
    {"code": "38", "name": "Ladakh"},
    {"code": "97", "name": "Other Territory"},
]

_STATE_BY_CODE = {s["code"]: s["name"] for s in INDIA_STATES}

# ---------------------------------------------------------------------------
# Pydantic models (field names identical to TS)
# ---------------------------------------------------------------------------


class LineItem(BaseModel):
    employee_name: str = ""
    service_month: str = ""
    sac: str = DEFAULT_SAC
    billing_hours: float = 0
    rate_per_hour: float = 0


class Buyer(BaseModel):
    name: str = ""
    address: str = ""
    pan: str = ""
    gstn: str = ""
    state_code: str = ""
    state_name: str = ""
    shipping: str = ""
    shipping_address: str = ""


class Invoice(BaseModel):
    invoice_no: str = DEFAULT_INVOICE_NO
    invoice_date: str = ""
    po_no: str | None = None
    po_date: str | None = None
    footer_text: str = DEFAULT_FOOTER
    buyer: Buyer = Field(default_factory=Buyer)
    items: list[LineItem] = Field(default_factory=list)
    #: Column labels follow the assignment's billing unit — a Monthly-billed
    #: project must not print "Rate/Hour" (customers query invoices over less).
    qty_label: str = "Billing Hours"
    rate_label: str = "Rate/Hour (INR)"


#: PE billing_unit -> (qty column, rate column) on the Tax Invoice.
UNIT_LABELS: dict[str, tuple[str, str]] = {
    "Hourly": ("Billing Hours", "Rate/Hour (INR)"),
    "Daily": ("Billing Days", "Rate/Day (INR)"),
    "Monthly": ("Billed Qty (Months)", "Rate/Month (INR)"),
    # Yearly assignments are invoiced at rate/12 through the Monthly branch.
    "Yearly": ("Billed Qty (Months)", "Rate/Month (INR)"),
}


def invoice_unit_labels(db, invoice) -> tuple[str, str]:
    """Qty/Rate column labels resolved from the invoice's timesheet assignment.

    Shared by the PDF renderer AND the invoice detail API, so the on-screen
    "View Tax Invoice" and the downloaded PDF can never disagree about whether
    this project bills per hour, day, month or year. Resolved at render time
    (nothing is stored), which retroactively corrects invoices generated
    before billing-unit labelling existed. Any lookup failure falls back to
    the historical hourly wording — labels must never block an invoice.
    """
    try:
        ts_row = getattr(invoice, "timesheet", None)
        pe_id = getattr(ts_row, "project_employee_id", None) if ts_row else None
        if pe_id:
            from models import ProjectEmployee
            pe = db.get(ProjectEmployee, pe_id)
            unit = getattr(getattr(pe, "billing_unit", None), "value",
                           getattr(pe, "billing_unit", None))
            if unit in UNIT_LABELS:
                return UNIT_LABELS[unit]
    except Exception:
        pass
    return UNIT_LABELS["Hourly"]


class Totals(BaseModel):
    subtotal: float
    intra: bool
    cgst: float
    sgst: float
    igst: float
    total_gst: float
    total: float
    amount_in_words: str
    tax_in_words: str


# ---------------------------------------------------------------------------
# Core numeric / tax helpers
# ---------------------------------------------------------------------------


def num(v: Any) -> float:
    try:
        if v is None or v == "":
            return 0.0
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def line_amount(it: LineItem | dict) -> float:
    if isinstance(it, dict):
        hours = num(it.get("billing_hours"))
        rate = num(it.get("rate_per_hour"))
    else:
        hours = num(it.billing_hours)
        rate = num(it.rate_per_hour)
    return hours * rate  # NO rounding


def normalize_state_code(code: Any) -> str:
    digits = re.sub(r"\D", "", str(code or ""))
    return digits[:2] if digits else ""


def buyer_state_code(buyer: Buyer | dict) -> str:
    if isinstance(buyer, dict):
        sc = buyer.get("state_code", "")
        gstn = buyer.get("gstn", "") or buyer.get("gstin", "")
    else:
        sc = buyer.state_code
        gstn = buyer.gstn
    norm = normalize_state_code(sc)
    if len(norm) == 2:
        return norm
    gst = str(gstn or "").strip().upper()
    if len(gst) >= 2 and gst[:2].isdigit():
        return gst[:2]
    return norm


def is_intra_state(buyer: Buyer | dict) -> bool:
    return buyer_state_code(buyer) == SELLER_STATE_CODE


def format_inr(value: float) -> str:
    """Display currency: 'INR ' + en-IN 2dp (ASCII — ₹ tofu in PDF fonts)."""
    # Force en-IN grouping for values that Python's default may miss for < 100000
    # Use locale-style manual grouping for Indian system.
    sign = "-" if value < 0 else ""
    abs_v = abs(float(value))
    whole, frac = f"{abs_v:.2f}".split(".")
    if len(whole) <= 3:
        grouped = whole
    else:
        last3 = whole[-3:]
        rest = whole[:-3]
        parts: list[str] = []
        while rest:
            parts.append(rest[-2:])
            rest = rest[:-2]
        grouped = ",".join(reversed(parts)) + "," + last3
    return f"INR {sign}{grouped}.{frac}"


# ---------------------------------------------------------------------------
# Amount in words (Indian Crore/Lakh — whole rupees only)
# ---------------------------------------------------------------------------

_ONES = [
    "", "One", "Two", "Three", "Four", "Five", "Six", "Seven", "Eight", "Nine",
    "Ten", "Eleven", "Twelve", "Thirteen", "Fourteen", "Fifteen", "Sixteen",
    "Seventeen", "Eighteen", "Nineteen",
]
_TENS = ["", "", "Twenty", "Thirty", "Forty", "Fifty", "Sixty", "Seventy", "Eighty", "Ninety"]


def _two_digits(n: int) -> str:
    if n < 20:
        return _ONES[n]
    t, o = divmod(n, 10)
    return f"{_TENS[t]}{_ONES[o] and ' ' + _ONES[o] or ''}".strip()


def _three_digits(n: int) -> str:
    """Hundreds + 'and' + tens/ones when both present."""
    if n <= 0:
        return ""
    h, rest = divmod(n, 100)
    parts: list[str] = []
    if h:
        parts.append(f"{_ONES[h]} Hundred")
    if rest:
        if h:
            parts.append("and")
        parts.append(_two_digits(rest))
    return " ".join(parts)


def amount_in_words(amount: Any) -> str:
    """Round to whole rupees; ignore paise. 'Rupees … Only'."""
    try:
        val = float(amount)
    except (TypeError, ValueError):
        val = 0.0
    negative = val < 0
    rupees = int(round(abs(val)))  # whole rupees; paise ignored via round-to-int of absolute

    if rupees == 0:
        words = "Rupees Zero Only"
        return f"Minus {words}" if negative else words

    crore, rem = divmod(rupees, 1_00_00_000)
    lakh, rem = divmod(rem, 1_00_000)
    thousand, rem = divmod(rem, 1_000)

    parts: list[str] = []
    if crore:
        # crore may exceed 99
        if crore < 100:
            parts.append(f"{_two_digits(crore)} Crore")
        else:
            # rare; spell with three-digit chunks simply
            parts.append(f"{_three_digits(crore % 1000)} Crore" if crore < 1000 else f"{crore} Crore")
    if lakh:
        parts.append(f"{_two_digits(lakh)} Lakh")
    if thousand:
        parts.append(f"{_two_digits(thousand)} Thousand")
    if rem:
        parts.append(_three_digits(rem))

    body = " ".join(p for p in parts if p).replace("  ", " ").strip()
    words = f"Rupees {body} Only"
    return f"Minus {words}" if negative else words


# ---------------------------------------------------------------------------
# Totals
# ---------------------------------------------------------------------------


def compute_totals(inv: Invoice | dict) -> Totals:
    if isinstance(inv, dict):
        items = inv.get("items") or []
        buyer = inv.get("buyer") or {}
    else:
        items = inv.items
        buyer = inv.buyer

    subtotal = sum(line_amount(it) for it in items)
    intra = is_intra_state(buyer)
    cgst = subtotal * 0.09 if intra else 0.0
    sgst = subtotal * 0.09 if intra else 0.0
    igst = 0.0 if intra else subtotal * 0.18
    total_gst = cgst + sgst + igst
    total = subtotal + total_gst
    return Totals(
        subtotal=subtotal,
        intra=intra,
        cgst=cgst,
        sgst=sgst,
        igst=igst,
        total_gst=total_gst,
        total=total,
        amount_in_words=amount_in_words(total),
        tax_in_words=amount_in_words(total_gst),
    )


# ---------------------------------------------------------------------------
# Line description + invoice numbering
# ---------------------------------------------------------------------------

_CSS_PREFIX_RE = re.compile(r"^contract\s+staffing\s+service\b", re.IGNORECASE)

_MONTH_ABBRS = ("Jan", "Feb", "Mar", "Apr", "May", "Jun",
                "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")
_MONTH_FULL = (
    "january", "february", "march", "april", "may", "june",
    "july", "august", "september", "october", "november", "december",
)

_MONTH_IN_DESC_RE = re.compile(
    r"\b(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*[\-' ]\s*(\d{4}|\d{2})\b",
    re.IGNORECASE,
)


def _normalize_month_label(label: str | None) -> str:
    """Normalize 'Jul' / 'July' / 'Jul 2026' / 'Jul-2026' → 'Jul' or 'Jul 2026'."""
    raw = (label or "").strip()
    if not raw:
        return ""
    # Prefer full month+year match first (handles July 2026 / Jul-2026).
    m = _MONTH_IN_DESC_RE.search(raw)
    if m:
        mon = m.group(1)[:3].title()
        yr = m.group(2)
        if len(yr) == 2:
            yr = f"20{yr}"
        return f"{mon} {yr}"
    key = re.split(r"[\s\-–—]+", raw, maxsplit=1)[0].lower()
    abbr = next((a for a in _MONTH_ABBRS if a.lower() == key), "")
    if not abbr:
        try:
            abbr = _MONTH_ABBRS[_MONTH_FULL.index(key)]
        except ValueError:
            return ""
    return abbr


def line_description(employee_name: str | None, month_label: str | None = None) -> str:
    """Canonical line text: "Contract Staffing Service <Name>[ - <Mon YYYY>]".

    `month_label` (e.g. "Jul 2026" from service_month_label) is appended when
    given and not already present in the text. Single-arg callers unchanged.
    """
    raw = (employee_name or "").strip()
    mon = _normalize_month_label(month_label)

    name = raw
    if not name:
        name = ""
    elif _CSS_PREFIX_RE.match(name):
        name = re.sub(
            r"^contract\s+staffing\s+service\s*", "", name, flags=re.I,
        ).strip()
        found = month_from_line_description(name)
        if found:
            mon = mon or found
            name = _MONTH_IN_DESC_RE.sub("", name).strip(" -–—'").strip()
        # "Name - Mon YYYY" after CSS strip
        m = re.match(
            r"^(.+?)\s*[-–—]\s*([A-Za-z]+)(?:\s+(\d{4}))?\s*$", name,
        )
        if m:
            name = m.group(1).strip()
            if not mon:
                mon = _normalize_month_label(
                    f"{m.group(2)} {m.group(3)}" if m.group(3) else m.group(2)
                )
    else:
        found_m = month_from_line_description(name)
        found_n = _employee_from_line_description(name)
        is_legacy = (
            " — " in name
            or " – " in name
            or "professional services" in name.lower()
            or (found_m and found_n != name)
        )
        if is_legacy:
            name = found_n
            mon = mon or found_m

    name = (name or "").strip()
    if name and mon:
        return f"Contract Staffing Service {name} - {mon}"
    if name:
        return f"Contract Staffing Service {name}"
    if mon:
        return f"Contract Staffing Service - {mon}"
    return "Contract Staffing Service"


def service_month_label(month: int | None, year: int | None) -> str:
    """Service-period label "Jul 2026" for line descriptions ("" when unknown)."""
    try:
        m = int(month or 0)
        y = int(year or 0)
    except (TypeError, ValueError):
        return ""
    if not (1 <= m <= 12) or y <= 0:
        return ""
    return f"{_MONTH_ABBRS[m - 1]} {y}"


def month_from_line_description(desc: str | None) -> str:
    """Extract a "Mon YYYY" service-period label already present in a line."""
    m = _MONTH_IN_DESC_RE.search(desc or "")
    if not m:
        return ""
    return _normalize_month_label(m.group(0))


def employee_from_line_description(desc: str | None) -> str:
    """Public alias: employee name from a stored line description (any month
    suffix stripped so line_description(emp, mon) never duplicates it)."""
    name = _employee_from_line_description(desc)
    return _MONTH_IN_DESC_RE.sub("", name).strip(" -–—'").strip()


_NO_WITH_SUFFIX = re.compile(r"^(.*-)(\d+)(-[A-Za-z]+)$")
_NO_PLAIN = re.compile(r"^(.*-)(\d+)$")


def invoice_no_at(base: str, offset: int) -> str:
    """Increment the numeric group before optional trailing letter suffix."""
    base = (base or DEFAULT_INVOICE_NO).strip() or DEFAULT_INVOICE_NO
    offset = int(offset or 0)
    m = _NO_WITH_SUFFIX.match(base)
    if m:
        prefix, digits, suffix = m.group(1), m.group(2), m.group(3)
        n = int(digits) + offset
        return f"{prefix}{str(n).zfill(len(digits))}{suffix}"
    m = _NO_PLAIN.match(base)
    if m:
        prefix, digits = m.group(1), m.group(2)
        n = int(digits) + offset
        return f"{prefix}{str(n).zfill(len(digits))}"
    return f"{base}-{offset}"


def sanitize_filename_part(text: str, cap: int = 80) -> str:
    s = re.sub(r"[^\w\-]+", "_", str(text or ""), flags=re.UNICODE)
    s = re.sub(r"_+", "_", s).strip("_")
    return (s or "invoice")[:cap]


def pdf_filename(inv: Invoice | dict) -> str:
    if isinstance(inv, dict):
        items = inv.get("items") or []
        no = inv.get("invoice_no") or DEFAULT_INVOICE_NO
        emp = (items[0].get("employee_name") if items and isinstance(items[0], dict)
               else (items[0].employee_name if items else "")) or ""
    else:
        no = inv.invoice_no or DEFAULT_INVOICE_NO
        emp = inv.items[0].employee_name if inv.items else ""
    return f"{sanitize_filename_part(emp)}_{sanitize_filename_part(no)}.pdf"


# ---------------------------------------------------------------------------
# Date helpers
# ---------------------------------------------------------------------------


def today_ddmmyyyy() -> str:
    return date.today().strftime("%d/%m/%Y")


def format_date_en_in(d: date | datetime | str | None) -> str:
    if d is None or d == "":
        return ""
    if isinstance(d, datetime):
        d = d.date()
    if isinstance(d, date):
        return d.strftime("%d/%m/%Y")
    s = str(d).strip()
    # Already dd/mm/yyyy?
    if re.match(r"^\d{2}/\d{2}/\d{4}$", s):
        return s
    # ISO
    try:
        if "T" in s:
            return datetime.fromisoformat(s.replace("Z", "+00:00")).date().strftime("%d/%m/%Y")
        return date.fromisoformat(s[:10]).strftime("%d/%m/%Y")
    except ValueError:
        return s


def excel_serial_to_date(serial: Any) -> str:
    """Excel serial → dd/mm/yyyy (base 1899-12-30)."""
    try:
        n = float(serial)
    except (TypeError, ValueError):
        return str(serial or "").strip()
    if n <= 0:
        return ""
    dt = datetime(1899, 12, 30) + timedelta(days=n)
    return dt.strftime("%d/%m/%Y")


# ---------------------------------------------------------------------------
# Excel import / template / summary
# ---------------------------------------------------------------------------

_HEADER_FILL = PatternFill("solid", fgColor="263E8F")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _norm_header(h: Any) -> str:
    s = str(h or "").lower().replace("\xa0", " ")
    s = re.sub(r"[:._\-()/]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def _find_col(headers: list[str], *aliases: str) -> int | None:
    norms = [_norm_header(h) for h in headers]
    alias_norms = [_norm_header(a) for a in aliases]
    for i, h in enumerate(norms):
        if h in alias_norms:
            return i
    # partial / contains
    for i, h in enumerate(norms):
        for a in alias_norms:
            if a and (a in h or h in a):
                return i
    return None


def _sheet_by_names(wb, names: Iterable[str], fallback_index: int):
    lower = {s.title.lower(): s for s in wb.worksheets}
    for n in names:
        if n.lower() in lower:
            return lower[n.lower()]
    if fallback_index < len(wb.worksheets):
        return wb.worksheets[fallback_index]
    return wb.worksheets[0]


def _row_values(row) -> list[Any]:
    return [c.value for c in row]


def parse_excel_bulk(data: bytes, *, base_invoice_no: str = DEFAULT_INVOICE_NO,
                     invoice_date: str | None = None) -> list[Invoice]:
    """Parse multi-buyer workbook → list of Invoice (recomputed amounts)."""
    wb = load_workbook(io.BytesIO(data), data_only=True)
    buyers_ws = _sheet_by_names(
        wb, ["Sheet1", "Buyer", "Buyers", "Buyer Details"], 0,
    )
    items_ws = _sheet_by_names(
        wb, ["Sheet2", "Line Items", "Items", "Services", "Line Item"], 1,
    )

    buyer_rows = list(buyers_ws.iter_rows(values_only=False))
    if not buyer_rows:
        return []
    b_headers = [c.value for c in buyer_rows[0]]
    c_ship = _find_col(b_headers, "Shipping Details", "Shipping", "Ship To", "Ship-To")
    c_addr1 = _find_col(b_headers, "Address1", "Address 1", "Shipping Address")
    c_pan = _find_col(b_headers, "PAN", "PAN No", "PAN Number")
    c_gstn = _find_col(b_headers, "GSTN ID", "GSTIN", "GSTN", "GSTIN ID", "GST No")
    c_scode = _find_col(b_headers, "State Code", "StateCode")
    c_sname = _find_col(b_headers, "State Name", "State")
    c_buyer = _find_col(b_headers, "Buyer Details", "Buyer", "Bill To", "Bill-To")
    c_addr2 = _find_col(b_headers, "Address2", "Address 2", "Billing Address")

    buyers: list[Buyer] = []
    for row in buyer_rows[1:]:
        vals = _row_values(row)
        if not any(v not in (None, "") for v in vals):
            continue

        def cell(idx: int | None) -> str:
            if idx is None or idx >= len(vals) or vals[idx] is None:
                return ""
            return str(vals[idx]).strip()

        gstn = cell(c_gstn).upper()
        scode = normalize_state_code(cell(c_scode))
        if len(scode) != 2 and len(gstn) >= 2 and gstn[:2].isdigit():
            scode = gstn[:2]
        sname = cell(c_sname)
        if not sname and scode in _STATE_BY_CODE:
            sname = _STATE_BY_CODE[scode]
        buyers.append(Buyer(
            name=cell(c_buyer),
            address=cell(c_addr2),
            pan=cell(c_pan),
            gstn=gstn,
            state_code=scode,
            state_name=sname,
            shipping=cell(c_ship),
            shipping_address=cell(c_addr1),
        ))

    item_rows = list(items_ws.iter_rows(values_only=False))
    items: list[LineItem] = []
    po_no = None
    po_date = None
    if item_rows:
        i_headers = [c.value for c in item_rows[0]]
        c_desc = _find_col(
            i_headers, "Description Of service", "Description of Services",
            "Description", "Employee Name", "Service",
        )
        c_sac = _find_col(i_headers, "SAC Code", "SAC", "HSN", "HSN/SAC")
        c_hrs = _find_col(
            i_headers, "Billing Hours", "Billable Hours", "Hours", "Billing Hour",
        )
        c_rate = _find_col(
            i_headers, "Rate/Hour", "Rate Per Hour", "Rate", "Rate Hour",
        )
        c_po = _find_col(
            i_headers, "P.O. No", "PO No", "PO Number", "P O No", "Purchase Order",
        )
        c_pod = _find_col(
            i_headers, "P.O. Date", "PO Date", "P O Date", "Purchase Order Date",
        )
        for row in item_rows[1:]:
            vals = _row_values(row)

            def cell(idx: int | None) -> Any:
                if idx is None or idx >= len(vals):
                    return None
                return vals[idx]

            raw_desc = str(cell(c_desc) or "").strip()
            emp = re.sub(
                r"^contract\s+staffing\s+service\s*", "", raw_desc, flags=re.I,
            ).strip()
            sac = str(cell(c_sac) or DEFAULT_SAC).strip() or DEFAULT_SAC
            hours = num(cell(c_hrs))
            rate = num(cell(c_rate))
            if not (emp or hours or rate or (sac and sac != DEFAULT_SAC) or raw_desc):
                continue
            items.append(LineItem(
                employee_name=emp or raw_desc,
                sac=sac,
                billing_hours=hours,
                rate_per_hour=rate,
            ))
            if po_no is None and cell(c_po) not in (None, ""):
                po_no = str(cell(c_po)).strip()
            if po_date is None and cell(c_pod) not in (None, ""):
                raw_pod = cell(c_pod)
                if isinstance(raw_pod, (datetime, date)):
                    po_date = format_date_en_in(raw_pod)
                elif isinstance(raw_pod, (int, float)):
                    po_date = excel_serial_to_date(raw_pod)
                else:
                    po_date = str(raw_pod).strip()

    inv_date = invoice_date or today_ddmmyyyy()
    if not buyers:
        # Single invoice with items only
        return [Invoice(
            invoice_no=invoice_no_at(base_invoice_no, 0),
            invoice_date=inv_date,
            po_no=po_no,
            po_date=po_date,
            buyer=Buyer(),
            items=items,
        )] if items else []

    out: list[Invoice] = []
    for i, buyer in enumerate(buyers):
        out.append(Invoice(
            invoice_no=invoice_no_at(base_invoice_no, i),
            invoice_date=inv_date,
            po_no=po_no,
            po_date=po_date,
            buyer=buyer,
            items=list(items),
        ))
    return out


def build_template_xlsx() -> bytes:
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    h1 = [
        "Shipping Details", "Address1", "PAN", "GSTN ID",
        "State Code", "State Name", "Buyer Details", "Address2",
    ]
    ws1.append(h1)
    for cell in ws1[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    ws1.append([
        "Acme Shipping", "Ship Addr Line", "ABCDE1234F", "27ABCDE1234F1Z5",
        "27", "Maharashtra", "Acme Buyer Pvt Ltd", "Bill Addr Line",
    ])

    ws2 = wb.create_sheet("Sheet2")
    h2 = [
        "Description Of service", "SAC Code", "Billing Hours", "Rate/Hour",
        "Amount (Rs.)", "P.O. No", "P.O. Date",
    ]
    ws2.append(h2)
    for cell in ws2[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    ws2.append([
        "Contract Staffing Service Rahul Sharma", DEFAULT_SAC, 176, 650,
        114400, "PO-1001", "01/04/2026",
    ])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_summary_xlsx(invoices: list[Invoice], day: date | None = None) -> bytes:
    day = day or date.today()
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    headers = [
        "Sr No.", "Description", "Invoice No", "Sub Total",
        "C-GST", "S-GST", "I-GST", "Grand Total",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    for i, inv in enumerate(invoices, start=1):
        t = compute_totals(inv)
        desc = ""
        if inv.items:
            desc = line_description(inv.items[0].employee_name)
        ws.append([
            i,
            desc,
            inv.invoice_no,
            round(t.subtotal, 2),
            round(t.cgst, 2),
            round(t.sgst, 2),
            round(t.igst, 2),
            round(t.total, 2),
        ])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# CRM invoice → Tax Invoice model
# ---------------------------------------------------------------------------


def _employee_from_line_description(desc: str | None) -> str:
    text = (desc or "").strip()
    if not text:
        return ""
    if _CSS_PREFIX_RE.match(text):
        return re.sub(
            r"^contract\s+staffing\s+service\s*", "", text, flags=re.I,
        ).strip()
    # "Name — project professional services, …"
    for sep in (" — ", " - ", " – "):
        if sep in text:
            return text.split(sep, 1)[0].strip()
    return text


def map_crm_invoice_to_tax_invoice(db, invoice) -> Invoice:
    """Map a stored CRM Invoice ORM row → Tax Invoice model."""
    from sqlalchemy import select
    from models import (
        Customer,
        CustomerBranch,
        POProjectAllocation,
        Project,
        PurchaseOrder,
    )
    from services.finance import resolve_billing_branch, resolve_buyer_state_code

    po = invoice.po
    if po is None and invoice.po_id:
        po = db.get(PurchaseOrder, invoice.po_id)

    project = invoice.project or db.get(Project, invoice.project_id)
    customer_id = po.customer_id if po else (project.customer_id if project else None)
    customer = db.get(Customer, customer_id) if customer_id else None

    billing_branch = None
    delivery_branch = None
    sac = DEFAULT_SAC
    if po is not None:
        if po.delivery_branch_id:
            delivery_branch = db.get(CustomerBranch, po.delivery_branch_id)
        alloc = db.execute(
            select(POProjectAllocation).where(
                POProjectAllocation.po_id == po.id,
                POProjectAllocation.project_id == invoice.project_id,
            )
        ).scalar_one_or_none()
        if alloc is not None and alloc.hsn_sac:
            sac = alloc.hsn_sac
    billing_branch = resolve_billing_branch(
        db, po=po, project=project, customer_id=customer_id,
    )
    if delivery_branch is None:
        delivery_branch = billing_branch

    def _addr(branch: CustomerBranch | None, *, delivery: bool) -> str:
        if branch is None:
            return ""
        parts: list[str] = []
        if delivery and branch.delivery_address:
            parts.append(branch.delivery_address.strip())
        else:
            if branch.billing_address:
                parts.append(branch.billing_address.strip())
            if branch.address_line_2:
                parts.append(branch.address_line_2.strip())
        city_line = ", ".join(p for p in [branch.city, branch.state, branch.pincode] if p)
        if city_line:
            parts.append(city_line)
        return ", ".join(parts)

    gstn = (billing_branch.gstin if billing_branch else "") or ""
    scode, _src = resolve_buyer_state_code(
        override=getattr(invoice, "buyer_state_code", None),
        branch=billing_branch,
    )
    sname = (billing_branch.state if billing_branch else "") or _STATE_BY_CODE.get(scode, "")
    if scode in _STATE_BY_CODE:
        sname = sname or _STATE_BY_CODE[scode]

    buyer_name = ""
    if billing_branch is not None:
        buyer_name = billing_branch.branch_legal_name or billing_branch.branch_name or ""
    if not buyer_name and customer is not None:
        buyer_name = customer.legal_entity_name or customer.name or ""

    ship_name = ""
    if delivery_branch is not None:
        ship_name = delivery_branch.branch_legal_name or delivery_branch.branch_name or ""
    if not ship_name:
        ship_name = buyer_name

    # Column labels from the billing unit of the timesheet's assignment.
    qty_label, rate_label = invoice_unit_labels(db, invoice)

    items: list[LineItem] = []
    lines = list(invoice.lines or [])
    lines.sort(key=lambda l: l.s_no or 0)
    for line in lines:
        desc = line.description
        items.append(LineItem(
            employee_name=employee_from_line_description(desc),
            service_month=month_from_line_description(desc),
            sac=sac,
            billing_hours=num(line.qty),
            rate_per_hour=num(line.rate),
        ))

    inv_no = (invoice.invoice_number or "").strip() or DEFAULT_INVOICE_NO
    po_no = po.po_number if po else None
    po_date = None
    if po is not None:
        raw = po.received_date or po.start_date
        po_date = format_date_en_in(raw) if raw else None

    return Invoice(
        qty_label=qty_label,
        rate_label=rate_label,
        invoice_no=inv_no,
        invoice_date=format_date_en_in(invoice.invoice_date) or today_ddmmyyyy(),
        po_no=po_no,
        po_date=po_date,
        buyer=Buyer(
            name=buyer_name,
            address=_addr(billing_branch, delivery=False),
            pan=(billing_branch.pan if billing_branch else "") or "",
            gstn=gstn or "",
            state_code=scode,
            state_name=sname,
            shipping=ship_name,
            shipping_address=_addr(delivery_branch, delivery=True),
        ),
        items=items,
    )


# ---------------------------------------------------------------------------
# PDF / HTML rendering
# ---------------------------------------------------------------------------


def _file_uri(path: Path) -> str:
    return path.resolve().as_uri() if path.exists() else ""


def _esc(s: Any) -> str:
    return (
        str(s or "")
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def render_invoice_html(inv: Invoice, totals: Totals | None = None) -> str:
    t = totals or compute_totals(inv)
    logo = _file_uri(LOGO_PATH)
    seal = _file_uri(SEAL_PATH)

    rows_html = []
    items = inv.items or []
    for i, it in enumerate(items, start=1):
        amt = line_amount(it)
        desc = line_description(it.employee_name, getattr(it, "service_month", None) or None)
        rows_html.append(
            "<tr>"
            f"<td class='c'>{i}</td>"
            f"<td>{_esc(desc)}</td>"
            f"<td class='c'>{_esc(it.sac or DEFAULT_SAC)}</td>"
            f"<td class='r'>{_esc(format_inr(num(it.billing_hours)).replace('INR ', ''))}</td>"
            f"<td class='r'>{_esc(format_inr(num(it.rate_per_hour)))}</td>"
            f"<td class='r'>{_esc(format_inr(amt))}</td>"
            "</tr>"
        )
    while len(rows_html) < 5:
        rows_html.append(
            "<tr class='spacer'><td>&nbsp;</td><td></td><td></td><td></td><td></td><td></td></tr>"
        )

    seller_addr_html = "<br/>".join(_esc(line) for line in SELLER_ADDRESS.split("\n"))
    buyer_addr_html = "<br/>".join(_esc(line) for line in (inv.buyer.address or "").split("\n") if line)

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8"/>
<title>TAX INVOICE {_esc(inv.invoice_no)}</title>
<style>
@page {{ size: A4; margin: 0; }}
* {{ box-sizing: border-box; }}
html, body {{
  margin: 0; padding: 0;
  font-family: Inter, "Segoe UI", Helvetica, Arial, sans-serif;
  font-size: 9.5pt; color: #1F2937; background: #fff;
  -webkit-print-color-adjust: exact; print-color-adjust: exact;
}}
.page {{
  width: 210mm; min-height: 297mm; padding: 8mm 9mm 6mm;
  display: flex; flex-direction: column; gap: 5px;
}}
:root {{
  --navy: #173B7A; --orange: #F59E0B; --border: #D9E2EC;
  --ink: #1F2937; --muted: #64748B; --light: #F8FAFC;
}}
.header {{ display: flex; gap: 0; border: 1px solid var(--border); }}
.header-left {{ width: 52%; padding: 8px 10px; }}
.header-right {{
  width: 48%; padding: 8px 10px; border-left: 1px solid var(--border);
}}
.logo {{ height: 36px; margin-bottom: 4px; }}
.seller-name {{
  font-weight: 700; color: var(--navy); font-size: 10.5pt;
  text-transform: uppercase; margin-bottom: 2px;
}}
.seller-meta {{ color: var(--muted); font-size: 8.5pt; line-height: 1.35; margin-top: 2px; }}
.icon-line {{
  display: flex; align-items: flex-start; gap: 5px;
  color: var(--muted); font-size: 8pt; line-height: 1.35; margin-top: 2px;
}}
.icon-line svg {{
  flex-shrink: 0; width: 11px; height: 11px; margin-top: 1px; color: var(--navy);
}}
.tax-title {{
  font-family: "Playfair Display", Georgia, serif;
  font-size: 26px; font-weight: 700; letter-spacing: 0.06em;
  text-transform: uppercase; color: var(--navy); margin: 0 0 8px; text-align: left;
}}
.meta-body {{ }}
.meta-table {{ width: auto; border-collapse: collapse; font-size: 8.5pt; }}
.meta-table td {{ padding: 2px 0; }}
.meta-table .k {{ color: var(--muted); white-space: nowrap; padding-right: 2px; }}
.meta-table .sep {{ color: var(--muted); padding: 0 6px 0 2px; }}
.wordmark {{ color: var(--orange); font-weight: 700; font-size: 9pt; }}
.cards {{ display: flex; gap: 6px; }}
.card {{ flex: 1; border: 1px solid var(--border); }}
.card-h {{
  background: var(--navy); color: #fff; font-weight: 700;
  font-size: 8.5pt; letter-spacing: 0.04em; padding: 4px 8px;
}}
.card-b {{ padding: 6px 8px; font-size: 8.5pt; line-height: 1.4; min-height: 72px; }}
.card-b .muted {{ color: var(--muted); }}
table.services {{
  width: 100%; border-collapse: collapse; border: 1px solid var(--border);
  table-layout: fixed;
}}
table.services th {{
  background: var(--navy); color: #fff; font-weight: 600;
  font-size: 8pt; padding: 5px 4px; text-align: left;
}}
table.services td {{
  border-top: 1px solid var(--border); padding: 4px; font-size: 8.5pt;
  vertical-align: top; height: 18px;
}}
table.services .c {{ text-align: center; }}
table.services .r {{ text-align: right; }}
table.services tr.spacer td {{ height: 16px; color: transparent; }}
.col-sno {{ width: 8%; }} .col-desc {{ width: 40%; }} .col-sac {{ width: 12%; }}
.col-hrs {{ width: 12%; }} .col-rate {{ width: 14%; }} .col-amt {{ width: 14%; }}
.gst-row {{ display: flex; gap: 6px; }}
.gst-box {{ flex: 1; border: 1px solid var(--border); }}
.gst-box table {{ width: 100%; border-collapse: collapse; font-size: 8.5pt; }}
.gst-box td {{ padding: 3px 8px; border-top: 1px solid var(--border); }}
.gst-box tr:first-child td {{ border-top: 0; }}
.gst-box .lab {{ color: var(--ink); }}
.gst-box .val {{ text-align: right; font-variant-numeric: tabular-nums; }}
.grand {{
  background: var(--navy); color: #fff; font-weight: 700;
}}
.grand .val {{ color: #fff; }}
.words {{
  border: 1px solid var(--border); padding: 6px 8px; font-size: 8.5pt;
  background: var(--light);
}}
.words .rs {{
  display: inline-block; width: 18px; height: 18px; line-height: 18px;
  text-align: center; background: var(--navy); color: #fff;
  border-radius: 3px; font-weight: 700; margin-right: 6px;
}}
.footer2 {{ display: flex; gap: 6px; }}
.footer2 > div {{
  flex: 1; border: 1px solid var(--border); padding: 6px 8px; font-size: 8pt;
}}
.footer2 h4 {{
  margin: 0 0 4px; color: var(--navy); font-size: 8.5pt; letter-spacing: 0.03em;
}}
.seal {{ height: 48px; margin-top: 4px; }}
.contact {{
  margin-top: auto; background: var(--navy); color: #fff;
  text-align: center; padding: 5px 8px; font-size: 8pt; letter-spacing: 0.02em;
}}
.contact span {{ margin: 0 8px; }}
</style>
<link rel="preconnect" href="https://fonts.googleapis.com"/>
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;600;700&family=Playfair+Display:wght@700&display=swap" rel="stylesheet"/>
</head>
<body>
<div class="page">
  <div class="header">
    <div class="header-left">
      {"<img class='logo' src='" + logo + "' alt='Karnex'/>" if logo else "<div class='wordmark'>KARNEX</div>"}
      <div class="seller-name">{_esc(SELLER_NAME)}</div>
      <div class="icon-line">
        <svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2"><path d="M20 10c0 6-8 12-8 12s-8-6-8-12a8 8 0 0 1 16 0Z"/><circle cx="12" cy="10" r="3"/></svg>
        <span>{seller_addr_html}</span>
      </div>
      <div class="icon-line">
        <svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2"><path d="M6 22V4a2 2 0 0 1 2-2h8a2 2 0 0 1 2 2v18Z"/><path d="M6 12H4a2 2 0 0 0-2 2v6a2 2 0 0 0 2 2h2"/><path d="M18 9h2a2 2 0 0 1 2 2v9a2 2 0 0 1-2 2h-2"/><path d="M10 6h4"/><path d="M10 10h4"/><path d="M10 14h4"/><path d="M10 18h4"/></svg>
        <span>State Name: {_esc(SELLER_STATE_NAME)} &nbsp; State Code: {_esc(SELLER_STATE_CODE)}</span>
      </div>
      <div class="icon-line">
        <svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2"><rect width="20" height="16" x="2" y="4" rx="2"/><path d="m22 7-8.97 5.7a1.94 1.94 0 0 1-2.06 0L2 7"/></svg>
        <span>Email: {_esc(SELLER_EMAIL)}</span>
      </div>
      <div class="icon-line">
        <svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2"><path d="M16 10h2"/><path d="M16 14h2"/><path d="M6.17 15a3 3 0 0 1 5.66 0"/><circle cx="9" cy="11" r="2"/><rect x="2" y="5" width="20" height="14" rx="2"/></svg>
        <span>CIN No.: {_esc(SELLER_CIN)}</span>
      </div>
    </div>
    <div class="header-right">
      <div class="tax-title">TAX INVOICE</div>
      <div class="meta-body">
        <table class="meta-table">
          <tr><td class="k">Invoice No.</td><td class="sep">:</td><td>{_esc(inv.invoice_no)}</td></tr>
          <tr><td class="k">Invoice Date</td><td class="sep">:</td><td>{_esc(inv.invoice_date)}</td></tr>
          <tr><td class="k">P.O. No.</td><td class="sep">:</td><td>{_esc(inv.po_no or "—")}</td></tr>
          <tr><td class="k">P.O. Date</td><td class="sep">:</td><td>{_esc(inv.po_date or "—")}</td></tr>
        </table>
      </div>
    </div>
  </div>

  <div class="cards">
    <div class="card">
      <div class="card-h">BUYER DETAILS</div>
      <div class="card-b">
        <strong>{_esc(inv.buyer.name)}</strong><br/>
        {buyer_addr_html or "—"}<br/>
        <span class="muted">PAN:</span> {_esc(inv.buyer.pan or "—")} &nbsp;
        <span class="muted">GSTIN:</span> {_esc(inv.buyer.gstn or "—")}<br/>
        <span class="muted">State Code:</span> {_esc(inv.buyer.state_code or "—")} &nbsp;
        <span class="muted">State Name:</span> {_esc(inv.buyer.state_name or "—")}<br/>
        <span class="muted">Shipping:</span> {_esc(inv.buyer.shipping or "—")}<br/>
        {_esc(inv.buyer.shipping_address or "")}
      </div>
    </div>
    <div class="card">
      <div class="card-h">SUPPLIER / BANK DETAILS</div>
      <div class="card-b">
        <span class="muted">PAN No.:</span> {_esc(SELLER_PAN)}<br/>
        <span class="muted">GSTIN No.:</span> {_esc(SELLER_GSTIN)}<br/>
        <span class="muted">Bank Name &amp; Address:</span> {_esc(BANK_NAME)}<br/>
        <span class="muted">Bank Account No.:</span> {_esc(BANK_ACC)}<br/>
        <span class="muted">IFSC Code:</span> {_esc(BANK_IFSC)}
      </div>
    </div>
  </div>

  <table class="services">
    <thead>
      <tr>
        <th class="col-sno">S. No.</th>
        <th class="col-desc">Description of Services</th>
        <th class="col-sac">SAC Code</th>
        <th class="col-hrs">{_esc(inv.qty_label)}</th>
        <th class="col-rate">{_esc(inv.rate_label)}</th>
        <th class="col-amt">Amount (INR)</th>
      </tr>
    </thead>
    <tbody>
      {"".join(rows_html)}
    </tbody>
  </table>

  <div class="gst-row">
    <div class="gst-box">
      <table>
        <tr><td class="lab">CGST @ 9%</td><td class="val">{_esc(format_inr(t.cgst))}</td></tr>
        <tr><td class="lab">SGST @ 9%</td><td class="val">{_esc(format_inr(t.sgst))}</td></tr>
        <tr><td class="lab">IGST @ 18%</td><td class="val">{_esc(format_inr(t.igst))}</td></tr>
        <tr><td class="lab">Total GST</td><td class="val">{_esc(format_inr(t.total_gst))}</td></tr>
      </table>
    </div>
    <div class="gst-box">
      <table>
        <tr><td class="lab">Sub Total</td><td class="val">{_esc(format_inr(t.subtotal))}</td></tr>
        <tr><td class="lab">CGST @ 9%</td><td class="val">{_esc(format_inr(t.cgst))}</td></tr>
        <tr><td class="lab">SGST @ 9%</td><td class="val">{_esc(format_inr(t.sgst))}</td></tr>
        <tr><td class="lab">IGST @ 18%</td><td class="val">{_esc(format_inr(t.igst))}</td></tr>
        <tr><td class="lab">Total GST Tax</td><td class="val">{_esc(format_inr(t.total_gst))}</td></tr>
        <tr class="grand"><td class="lab">GRAND TOTAL</td><td class="val">{_esc(format_inr(t.total))}</td></tr>
      </table>
    </div>
  </div>

  <div class="words">
    <div><strong>Amount Chargeable (in Words):</strong> {_esc(t.amount_in_words)}</div>
    <div style="margin-top:3px"><strong>Tax Amount (in Words):</strong> {_esc(t.tax_in_words)}</div>
  </div>

  <div class="footer2">
    <div>
      <h4>BANK DETAILS</h4>
      Bank Name: {_esc(BANK_NAME_SHORT)}<br/>
      Account No.: {_esc(BANK_ACC)}<br/>
      IFSC Code: {_esc(BANK_IFSC)}<br/>
      Branch: {_esc(BANK_BRANCH)}
    </div>
    <div>
      <h4>Declaration</h4>
      {_esc(inv.footer_text)}<br/><br/>
      <strong>For Karnex Software Solutions Pvt. Ltd.</strong><br/>
      {"<img class='seal' src='" + seal + "' alt='seal'/>" if seal else ""}
      <div>Authorized Signatory</div>
    </div>
  </div>

  <div class="contact">
    <span>{_esc(CONTACT_WEB)}</span>·
    <span>{_esc(CONTACT_EMAIL)}</span>·
    <span>{_esc(CONTACT_PHONE)}</span>·
    <span>{_esc(CONTACT_LOC)}</span>
  </div>
</div>
</body>
</html>"""


def _weasyprint_available() -> bool:
    try:
        import weasyprint  # noqa: F401
        # Import may succeed but GTK missing — probe HTML()
        from weasyprint import HTML  # noqa: F401
        return True
    except Exception:
        return False


def _pdf_via_weasyprint(html: str) -> bytes:
    from weasyprint import HTML
    return HTML(string=html, base_url=str(STATIC_DIR)).write_pdf()


def _pdf_via_reportlab(inv: Invoice, totals: Totals) -> bytes:
    """A4 PDF fallback when WeasyPrint/GTK is unavailable (Windows)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib.colors import HexColor, white, black
    from reportlab.platypus import (
        SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image,
    )
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_LEFT

    navy = HexColor("#173B7A")
    border = HexColor("#D9E2EC")
    muted = HexColor("#64748B")

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=9 * mm, rightMargin=9 * mm,
        topMargin=8 * mm, bottomMargin=6 * mm,
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="InvTitle", fontName="Times-Bold", fontSize=18,
                              textColor=navy, alignment=TA_LEFT, spaceAfter=4))
    styles.add(ParagraphStyle(name="InvSmall", fontSize=7.5, textColor=muted, leading=10))
    styles.add(ParagraphStyle(name="InvBody", fontSize=8, textColor=black, leading=11))
    styles.add(ParagraphStyle(name="InvHead", fontSize=9, textColor=navy, fontName="Helvetica-Bold"))
    styles.add(ParagraphStyle(name="InvWhite", fontSize=8, textColor=white, fontName="Helvetica-Bold"))

    story: list[Any] = []
    logo_flow = None
    if LOGO_PATH.exists():
        try:
            logo_flow = Image(str(LOGO_PATH), width=90, height=28)
        except Exception:
            logo_flow = None

    left_bits = []
    if logo_flow:
        left_bits.append(logo_flow)
    left_bits.append(Paragraph(f"<b>{_esc(SELLER_NAME)}</b>", styles["InvHead"]))
    left_bits.append(Paragraph(SELLER_ADDRESS.replace("\n", "<br/>"), styles["InvSmall"]))
    left_bits.append(Paragraph(
        f"Email: {_esc(SELLER_EMAIL)}<br/>State: {_esc(SELLER_STATE_NAME)} ({_esc(SELLER_STATE_CODE)})"
        f"<br/>CIN: {_esc(SELLER_CIN)}",
        styles["InvSmall"],
    ))

    right_bits = [
        Paragraph("TAX INVOICE", styles["InvTitle"]),
        Paragraph(
            f"Invoice No.: <b>{_esc(inv.invoice_no)}</b><br/>"
            f"Invoice Date: {_esc(inv.invoice_date)}<br/>"
            f"P.O. No.: {_esc(inv.po_no or '—')}<br/>"
            f"P.O. Date: {_esc(inv.po_date or '—')}<br/>"
            f"Supplier PAN No. {_esc(SELLER_PAN)}<br/>"
            f"GSTIN No. {_esc(SELLER_GSTIN)}",
            styles["InvBody"],
        ),
    ]

    # Build nested tables for header columns
    left_tbl = Table([[b] for b in left_bits], colWidths=[95 * mm])
    right_tbl = Table([[b] for b in right_bits], colWidths=[88 * mm])
    header = Table([[left_tbl, right_tbl]], colWidths=[98 * mm, 90 * mm])
    header.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 0.5, border),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("LINEBEFORE", (1, 0), (1, 0), 0.5, border),
    ]))
    story.append(header)
    story.append(Spacer(1, 4))

    buyer_p = Paragraph(
        f"<b>{_esc(inv.buyer.name)}</b><br/>{_esc(inv.buyer.address)}<br/>"
        f"PAN: {_esc(inv.buyer.pan or '—')} &nbsp; GSTIN: {_esc(inv.buyer.gstn or '—')}<br/>"
        f"State Code: {_esc(inv.buyer.state_code or '—')} &nbsp; "
        f"State Name: {_esc(inv.buyer.state_name or '—')}",
        styles["InvBody"],
    )
    bank_p = Paragraph(
        f"PAN No.: {_esc(SELLER_PAN)}<br/>GSTIN No.: {_esc(SELLER_GSTIN)}<br/>"
        f"Bank Name &amp; Address: {_esc(BANK_NAME)}<br/>"
        f"Bank Account No.: {_esc(BANK_ACC)}<br/>IFSC Code: {_esc(BANK_IFSC)}",
        styles["InvBody"],
    )
    bh = Paragraph("BUYER DETAILS", styles["InvWhite"])
    sh = Paragraph("SUPPLIER / BANK DETAILS", styles["InvWhite"])
    cards = Table(
        [[bh, sh], [buyer_p, bank_p]],
        colWidths=[94 * mm, 94 * mm],
    )
    cards.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), navy),
        ("BOX", (0, 0), (0, -1), 0.5, border),
        ("BOX", (1, 0), (1, -1), 0.5, border),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, 0), 4),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 4),
    ]))
    story.append(cards)
    story.append(Spacer(1, 4))

    svc_data = [[
        "S. No.", "Description of Services", "SAC Code",
        inv.qty_label, inv.rate_label, "Amount (INR)",
    ]]
    for i, it in enumerate(inv.items or [], start=1):
        svc_data.append([
            str(i),
            line_description(it.employee_name, getattr(it, "service_month", None) or None),
            it.sac or DEFAULT_SAC,
            f"{num(it.billing_hours):,.2f}",
            format_inr(num(it.rate_per_hour)),
            format_inr(line_amount(it)),
        ])
    while len(svc_data) < 6:
        svc_data.append(["", "", "", "", "", ""])

    col_w = [15 * mm, 72 * mm, 22 * mm, 22 * mm, 28 * mm, 29 * mm]
    svc = Table(svc_data, colWidths=col_w, repeatRows=1)
    svc.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), navy),
        ("TEXTCOLOR", (0, 0), (-1, 0), white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("GRID", (0, 0), (-1, -1), 0.4, border),
        ("ALIGN", (0, 0), (0, -1), "CENTER"),
        ("ALIGN", (3, 1), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(svc)
    story.append(Spacer(1, 4))

    left_gst = [
        ["CGST @ 9%", format_inr(totals.cgst)],
        ["SGST @ 9%", format_inr(totals.sgst)],
        ["IGST @ 18%", format_inr(totals.igst)],
        ["Total GST", format_inr(totals.total_gst)],
    ]
    right_gst = [
        ["Sub Total", format_inr(totals.subtotal)],
        ["CGST @ 9%", format_inr(totals.cgst)],
        ["SGST @ 9%", format_inr(totals.sgst)],
        ["IGST @ 18%", format_inr(totals.igst)],
        ["Total GST Tax", format_inr(totals.total_gst)],
        ["GRAND TOTAL", format_inr(totals.total)],
    ]
    lg = Table(left_gst, colWidths=[50 * mm, 40 * mm])
    rg = Table(right_gst, colWidths=[50 * mm, 40 * mm])
    for tbl in (lg, rg):
        tbl.setStyle(TableStyle([
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("ALIGN", (1, 0), (1, -1), "RIGHT"),
            ("GRID", (0, 0), (-1, -1), 0.4, border),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ]))
    rg.setStyle(TableStyle([
        ("BACKGROUND", (0, -1), (-1, -1), navy),
        ("TEXTCOLOR", (0, -1), (-1, -1), white),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
    ]))
    gst_row = Table([[lg, rg]], colWidths=[94 * mm, 94 * mm])
    story.append(gst_row)
    story.append(Spacer(1, 4))

    story.append(Paragraph(
        f"<b>Amount Chargeable (in Words):</b> {_esc(totals.amount_in_words)}<br/>"
        f"<b>Tax Amount (in Words):</b> {_esc(totals.tax_in_words)}",
        styles["InvBody"],
    ))
    story.append(Spacer(1, 4))

    bank_f = Paragraph(
        f"<b>BANK DETAILS</b><br/>Bank Name: {_esc(BANK_NAME_SHORT)}<br/>"
        f"Account No.: {_esc(BANK_ACC)}<br/>IFSC Code: {_esc(BANK_IFSC)}<br/>"
        f"Branch: {_esc(BANK_BRANCH)}",
        styles["InvBody"],
    )
    decl_bits = [
        Paragraph(f"<b>Declaration</b><br/>{_esc(inv.footer_text)}", styles["InvBody"]),
        Paragraph("<b>For Karnex Software Solutions Pvt. Ltd.</b>", styles["InvBody"]),
    ]
    if SEAL_PATH.exists():
        try:
            decl_bits.append(Image(str(SEAL_PATH), width=40, height=40))
        except Exception:
            pass
    decl_bits.append(Paragraph("Authorized Signatory", styles["InvSmall"]))
    decl_tbl = Table([[b] for b in decl_bits], colWidths=[90 * mm])
    foot = Table([[bank_f, decl_tbl]], colWidths=[94 * mm, 94 * mm])
    foot.setStyle(TableStyle([
        ("BOX", (0, 0), (0, 0), 0.4, border),
        ("BOX", (1, 0), (1, 0), 0.4, border),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(foot)
    story.append(Spacer(1, 6))
    contact = Table([[
        f"{CONTACT_WEB}  ·  {CONTACT_EMAIL}  ·  {CONTACT_PHONE}  ·  {CONTACT_LOC}"
    ]], colWidths=[188 * mm])
    contact.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), navy),
        ("TEXTCOLOR", (0, 0), (-1, -1), white),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(contact)

    doc.build(story)
    return buf.getvalue()


@dataclass
class PdfResult:
    content: bytes
    media_type: str
    filename: str
    engine: str  # weasyprint | reportlab | html


def render_pdf(inv: Invoice) -> PdfResult:
    totals = compute_totals(inv)
    html = render_invoice_html(inv, totals)
    fname = pdf_filename(inv)
    try:
        pdf = _pdf_via_weasyprint(html)
        return PdfResult(pdf, "application/pdf", fname, "weasyprint")
    except Exception:
        try:
            pdf = _pdf_via_reportlab(inv, totals)
            return PdfResult(pdf, "application/pdf", fname, "reportlab")
        except Exception:
            return PdfResult(
                html.encode("utf-8"),
                "text/html; charset=utf-8",
                fname.replace(".pdf", ".html"),
                "html",
            )


def build_bulk_zip(invoices: list[Invoice], day: date | None = None) -> bytes:
    day = day or date.today()
    day_s = day.isoformat()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for inv in invoices:
            result = render_pdf(inv)
            zf.writestr(result.filename, result.content)
        summary = build_summary_xlsx(invoices, day)
        zf.writestr(f"Invoice_Summary_{day_s}.xlsx", summary)
    # Rename is handled by Content-Disposition on the response
    return buf.getvalue()


def seller_public_dict() -> dict:
    return {
        "name": SELLER_NAME,
        "address": SELLER_ADDRESS,
        "state_name": SELLER_STATE_NAME,
        "state_code": SELLER_STATE_CODE,
        "email": SELLER_EMAIL,
        "cin": SELLER_CIN,
        "pan": SELLER_PAN,
        "gstin": SELLER_GSTIN,
        "bank_name": BANK_NAME,
        "bank_name_short": BANK_NAME_SHORT,
        "bank_acc": BANK_ACC,
        "ifsc": BANK_IFSC,
        "branch": BANK_BRANCH,
        "contact_web": CONTACT_WEB,
        "contact_email": CONTACT_EMAIL,
        "contact_phone": CONTACT_PHONE,
        "contact_loc": CONTACT_LOC,
    }
